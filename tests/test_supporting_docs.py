"""supporting_docs.analyse and its pure helpers -- group reconciliation, date
validation, .docx/PDF text extraction.

The network seam is supporting_docs._create; every analyse() test fakes it,
same pattern as tests/test_brief_extraction.py. reconcile_group/
resolve_group_scope/extract_text need no seam at all.
"""
import io
import json
import types

import docx
import fitz
import pytest

import briefs
import supporting_docs


class FakeResp:
    def __init__(self, text, stop_reason="end_turn"):
        self.content = [types.SimpleNamespace(type="text", text=text)]
        self.stop_reason = stop_reason


@pytest.fixture(autouse=True)
def _isolate(monkeypatch):
    monkeypatch.setattr(supporting_docs.tagging, "get_client", lambda: object())


# --- group reconciliation ----------------------------------------------------


@pytest.mark.parametrize("group_text, cohort_group, expected", [
    (None, "gp3", None),
    ("", "gp3", None),
    ("Groups 1 & 2", "gp3", "other"),
    ("Groups 3 & 4", "gp3", "mine"),
    ("Group 3", "gp3", "mine"),
    ("Groups 1 & 2", None, "unresolved"),  # no cohort_group recorded at all
    ("All groups", "gp3", "unresolved"),  # no digits to compare on the doc's side
])
def test_reconcile_group(group_text, cohort_group, expected):
    assert supporting_docs.reconcile_group(group_text, cohort_group) == expected


def test_resolve_group_scope_proposes_only_the_users_group():
    # The exact scenario from the fault: two group-tagged dates, one per pair
    # of groups, and the user is in group 3.
    items = [
        {"title": "Bring toile A", "due_date": "2026-09-21", "group": "Groups 1 & 2"},
        {"title": "Bring toile B", "due_date": "2026-09-23", "group": "Groups 3 & 4"},
    ]
    kept = supporting_docs.resolve_group_scope(items, "gp3")
    assert [i["title"] for i in kept] == ["Bring toile B"]
    assert kept[0]["group_match"] == "mine"


def test_resolve_group_scope_keeps_unresolved_and_ungrouped_items():
    items = [
        {"title": "No group at all", "group": None},
        {"title": "Can't tell", "group": "All groups"},
        {"title": "Not mine", "group": "Groups 1 & 2"},
    ]
    kept = supporting_docs.resolve_group_scope(items, "gp3")
    titles = {i["title"] for i in kept}
    assert titles == {"No group at all", "Can't tell"}
    assert "group_match" not in next(i for i in kept if i["title"] == "No group at all")
    assert next(i for i in kept if i["title"] == "Can't tell")["group_match"] == "unresolved"


# --- text extraction ---------------------------------------------------------


def test_extract_text_reads_a_docx(tmp_path):
    path = tmp_path / "workshop.docx"
    doc = docx.Document()
    doc.add_paragraph("Construction workshop -- materials list.")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Fabric shears"
    table.rows[0].cells[1].text = "1 pair"
    doc.save(path)

    text = supporting_docs.extract_text(path)

    assert "Construction workshop" in text
    assert "Fabric shears" in text
    assert "1 pair" in text


def test_extract_text_reads_a_pdf(tmp_path):
    path = tmp_path / "reading-list.pdf"
    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_text((72, 100), "Reading list: bring three references.")
    pdf.save(path)
    pdf.close()

    assert "Reading list" in supporting_docs.extract_text(path)


# --- analyse() ----------------------------------------------------------------

EXTRACTION = json.dumps({
    "summary": "A workshop materials list, split by group.",
    "sessions": [
        {"label": "Construction workshop", "date": "2026-09-21", "group": "Groups 1 & 2"},
        {"label": "Construction workshop", "date": "2026-09-23", "group": "Groups 3 & 4"},
    ],
    "preparation_tasks": [
        {"title": "Bring fabric shears", "due_date": "2026-09-21",
         "group": "Groups 1 & 2", "session_index": 0},
        {"title": "Bring your first-year toile", "due_date": "2026-09-23",
         "group": "Groups 3 & 4", "session_index": 1},
        {"title": "Bring a metre of calico", "due_date": "2026-09-23",
         "group": "Groups 3 & 4", "session_index": 1},
    ],
})


def _fake_create(text=EXTRACTION, stop_reason="end_turn"):
    def create(client, prompt, max_tokens=4096):
        return FakeResp(text, stop_reason)
    return create


def test_analyse_proposes_only_the_users_group_and_marks_the_others_absent(monkeypatch):
    monkeypatch.setattr(supporting_docs, "_create", _fake_create())

    out = supporting_docs.analyse("a document", cohort_group="gp3")

    # Sessions stay whole (informational), the fabric-shears task for the
    # OTHER group is not proposed at all -- the group-split acceptance case.
    assert len(out["sessions"]) == 2
    titles = [t["title"] for t in out["preparation_tasks"]]
    assert titles == ["Bring your first-year toile", "Bring a metre of calico"]
    assert all(t["group_match"] == "mine" for t in out["preparation_tasks"])


def test_analyse_produces_one_task_per_material_never_one_combined_task(monkeypatch):
    # The materials-list fault: three distinct items must stay three distinct
    # tasks, each with its own due date -- never folded into "gather materials".
    monkeypatch.setattr(supporting_docs, "_create", _fake_create())
    out = supporting_docs.analyse("a document", cohort_group="gp4")
    assert len(out["preparation_tasks"]) == 2  # groups 3&4 kept, 1&2 dropped
    keys = {t["source_key"] for t in out["preparation_tasks"]}
    assert len(keys) == 2  # each got its own stable key


def test_analyse_marks_sessions_already_on_the_timetable(monkeypatch):
    monkeypatch.setattr(supporting_docs, "_create", _fake_create())
    out = supporting_docs.analyse(
        "a document", cohort_group="gp3", commitment_dates={"2026-09-21"},
    )
    by_date = {s["date"]: s for s in out["sessions"]}
    assert by_date["2026-09-21"]["matched_commitment"] is True
    assert by_date["2026-09-23"]["matched_commitment"] is False


def test_analyse_flags_a_date_outside_the_range_without_correcting_it(monkeypatch):
    monkeypatch.setattr(supporting_docs, "_create", _fake_create())
    ctx = briefs.date_context(("2026-09-21", "2026-10-27"))

    out = supporting_docs.analyse("a document", date_context=ctx, cohort_group="gp3")

    # Both remaining (group-3&4) tasks are due 2026-09-23, inside range.
    assert all("date_suspect" not in t for t in out["preparation_tasks"])

    far_extraction = EXTRACTION.replace("2026-09-23", "2025-01-05")
    monkeypatch.setattr(supporting_docs, "_create", _fake_create(far_extraction))
    out = supporting_docs.analyse("a document", date_context=ctx, cohort_group="gp3")
    assert all(t["due_date"] == "2025-01-05" for t in out["preparation_tasks"])
    assert all(t["date_suspect"]["checked_against"] == ["2026-09-21", "2026-10-27"]
               for t in out["preparation_tasks"])


def test_a_truncated_reply_raises(monkeypatch):
    monkeypatch.setattr(supporting_docs, "_create", _fake_create(stop_reason="max_tokens"))
    with pytest.raises(briefs.BriefExtractionError) as excinfo:
        supporting_docs.analyse("a document")
    assert excinfo.value.reason == "truncated"


def test_a_malformed_reply_raises(monkeypatch):
    monkeypatch.setattr(supporting_docs, "_create", _fake_create("not json{"))
    with pytest.raises(briefs.BriefExtractionError) as excinfo:
        supporting_docs.analyse("a document")
    assert excinfo.value.reason == "malformed"
